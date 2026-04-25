# export.py — Módulo de exportación para stock.app
# Agregar este archivo en la carpeta raíz del proyecto (junto a app.py)
# Luego en app.py agregar: from export import export_bp y app.register_blueprint(export_bp)

from flask import Blueprint, session, send_file
from functools import wraps
import sqlite3, io, json
from datetime import datetime

export_bp = Blueprint('export', __name__)
DATABASE = 'stockapp.db'

def get_db():
    db = sqlite3.connect(DATABASE)
    db.row_factory = sqlite3.Row
    return db

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'usuario_id' not in session:
            from flask import jsonify
            return jsonify({'error': 'No autenticado'}), 401
        return f(*args, **kwargs)
    return decorated

def perm_required(permiso):
    def decorator(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            if 'usuario_id' not in session:
                from flask import jsonify
                return jsonify({'error': 'No autenticado'}), 401
            perms = session.get('perms', {})
            if not perms.get(permiso):
                from flask import jsonify
                return jsonify({'error': 'Sin permiso'}), 403
            return f(*args, **kwargs)
        return decorated
    return decorator

def fmt_clp(n):
    return f"${int(n):,}".replace(',', '.')

def nombre_negocio():
    return "Mi negocio"

# ── EXCEL ─────────────────────────────────────────────────────────────────

def excel_response(wb, filename):
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

def estilo_excel(ws, titulo):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # Colores
    COLOR_HEADER = "0D1220"
    COLOR_ACCENT = "4F8EFF"
    COLOR_TEXT   = "E8F0FF"
    COLOR_SUBHEAD = "1E2D45"

    # Fila 1: Título del negocio
    ws.insert_rows(1)
    ws.insert_rows(1)
    ws['A1'] = nombre_negocio()
    ws['A1'].font = Font(name='Calibri', bold=True, size=14, color=COLOR_TEXT)
    ws['A1'].fill = PatternFill("solid", fgColor=COLOR_HEADER)
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 28

    ws['A2'] = titulo
    ws['A2'].font = Font(name='Calibri', bold=True, size=11, color="7A92B8")
    ws['A2'].fill = PatternFill("solid", fgColor=COLOR_HEADER)
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[2].height = 20

    # Estilo headers (fila 3)
    thin = Side(style='thin', color="1E2D45")
    border = Border(bottom=thin)
    for cell in ws[3]:
        cell.font = Font(name='Calibri', bold=True, size=10, color="7A92B8")
        cell.fill = PatternFill("solid", fgColor=COLOR_SUBHEAD)
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = border
    ws.row_dimensions[3].height = 22

    # Estilo filas de datos
    for row in ws.iter_rows(min_row=4):
        for i, cell in enumerate(row):
            cell.font = Font(name='Calibri', size=10, color="C8D8F0")
            cell.fill = PatternFill("solid", fgColor="0C1422" if row[0].row % 2 == 0 else "080B12")
            cell.alignment = Alignment(horizontal='left', vertical='center')
        ws.row_dimensions[row[0].row].height = 18

    # Ancho automático columnas
    for col in ws.columns:
        max_w = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_w = max(max_w, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_w + 4, 40)

@export_bp.route('/api/export/excel/inventario')
@login_required
def excel_inventario():
    from openpyxl import Workbook
    db = get_db()
    productos = db.execute('SELECT * FROM productos ORDER BY nombre').fetchall()
    db.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append(['Producto', 'Categoría', 'Stock', 'Stock mínimo', 'Precio costo', 'Precio venta', 'Estado', 'Valor stock'])
    for p in productos:
        estado = 'Agotado' if p['stock'] == 0 else ('Stock bajo' if p['stock'] <= p['stock_minimo'] else 'OK')
        ws.append([p['nombre'], p['categoria'], p['stock'], p['stock_minimo'],
                   p['precio_costo'], p['precio_venta'], estado, p['stock'] * p['precio_venta']])
    estilo_excel(ws, f"Inventario — {datetime.now().strftime('%d/%m/%Y')}")
    return excel_response(wb, f"inventario_{datetime.now().strftime('%Y%m%d')}.xlsx")

@export_bp.route('/api/export/excel/ventas')
@login_required
def excel_ventas():
    from openpyxl import Workbook
    db = get_db()
    ventas = db.execute(
        'SELECT v.*, c.nombre as cliente_nombre, u.nombre as usuario_nombre FROM ventas v '
        'LEFT JOIN clientes c ON v.cliente_id=c.id '
        'LEFT JOIN usuarios u ON v.usuario_id=u.id '
        'ORDER BY v.creado_en DESC').fetchall()
    db.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws.append(['Fecha', 'Hora', 'Usuario', 'Cliente', 'Método pago', 'Total'])
    for v in ventas:
        dt = datetime.fromisoformat(v['creado_en'])
        ws.append([dt.strftime('%d/%m/%Y'), dt.strftime('%H:%M'),
                   v['usuario_nombre'] or '—', v['cliente_nombre'] or 'Público general',
                   v['metodo_pago'], v['total']])
    estilo_excel(ws, f"Historial de ventas — {datetime.now().strftime('%d/%m/%Y')}")
    return excel_response(wb, f"ventas_{datetime.now().strftime('%Y%m%d')}.xlsx")

@export_bp.route('/api/export/excel/clientes')
@login_required
def excel_clientes():
    from openpyxl import Workbook
    db = get_db()
    clientes = db.execute('SELECT * FROM clientes ORDER BY nombre').fetchall()
    ventas = db.execute('SELECT cliente_id, SUM(total) as total, COUNT(*) as nventas FROM ventas GROUP BY cliente_id').fetchall()
    db.close()

    ventas_map = {v['cliente_id']: v for v in ventas}
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append(['Nombre', 'Teléfono', 'Email', 'Tipo', 'N° compras', 'Total gastado', 'Notas'])
    for c in clientes:
        vdata = ventas_map.get(c['id'])
        ws.append([c['nombre'], c['telefono'] or '—', c['email'] or '—', c['tipo'],
                   vdata['nventas'] if vdata else 0,
                   vdata['total'] if vdata else 0,
                   c['notas'] or ''])
    estilo_excel(ws, f"Clientes — {datetime.now().strftime('%d/%m/%Y')}")
    return excel_response(wb, f"clientes_{datetime.now().strftime('%Y%m%d')}.xlsx")

@export_bp.route('/api/export/excel/stock-bajo')
@login_required
def excel_stock_bajo():
    from openpyxl import Workbook
    db = get_db()
    productos = db.execute(
        'SELECT * FROM productos WHERE stock <= stock_minimo ORDER BY stock ASC').fetchall()
    db.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Stock bajo"
    ws.append(['Producto', 'Categoría', 'Stock actual', 'Stock mínimo', 'Estado', 'Precio costo'])
    for p in productos:
        estado = 'AGOTADO' if p['stock'] == 0 else 'STOCK BAJO'
        ws.append([p['nombre'], p['categoria'], p['stock'], p['stock_minimo'],
                   estado, p['precio_costo']])
    estilo_excel(ws, f"Productos para reponer — {datetime.now().strftime('%d/%m/%Y')}")
    return excel_response(wb, f"stock_bajo_{datetime.now().strftime('%Y%m%d')}.xlsx")

@export_bp.route('/api/export/excel/completo')
@perm_required('exportar')
def excel_completo():
    from openpyxl import Workbook
    db = get_db()
    productos = db.execute('SELECT * FROM productos ORDER BY nombre').fetchall()
    clientes  = db.execute('SELECT * FROM clientes ORDER BY nombre').fetchall()
    ventas    = db.execute(
        'SELECT v.*, c.nombre as cn, u.nombre as un FROM ventas v '
        'LEFT JOIN clientes c ON v.cliente_id=c.id '
        'LEFT JOIN usuarios u ON v.usuario_id=u.id '
        'ORDER BY v.creado_en DESC').fetchall()
    db.close()

    wb = Workbook()

    # Hoja inventario
    ws1 = wb.active
    ws1.title = "Inventario"
    ws1.append(['Producto', 'Categoría', 'Stock', 'Stock mínimo', 'P. Costo', 'P. Venta', 'Estado'])
    for p in productos:
        estado = 'Agotado' if p['stock'] == 0 else ('Stock bajo' if p['stock'] <= p['stock_minimo'] else 'OK')
        ws1.append([p['nombre'], p['categoria'], p['stock'], p['stock_minimo'],
                    p['precio_costo'], p['precio_venta'], estado])
    estilo_excel(ws1, "Inventario")

    # Hoja ventas
    ws2 = wb.create_sheet("Ventas")
    ws2.append(['Fecha', 'Cliente', 'Método pago', 'Total'])
    for v in ventas:
        dt = datetime.fromisoformat(v['creado_en'])
        ws2.append([dt.strftime('%d/%m/%Y %H:%M'), v['cn'] or 'General', v['metodo_pago'], v['total']])
    estilo_excel(ws2, "Historial de ventas")

    # Hoja clientes
    ws3 = wb.create_sheet("Clientes")
    ws3.append(['Nombre', 'Teléfono', 'Email', 'Tipo', 'Notas'])
    for c in clientes:
        ws3.append([c['nombre'], c['telefono'] or '—', c['email'] or '—', c['tipo'], c['notas'] or ''])
    estilo_excel(ws3, "Clientes")

    return excel_response(wb, f"respaldo_completo_{datetime.now().strftime('%Y%m%d')}.xlsx")

# ── PDF ───────────────────────────────────────────────────────────────────

def pdf_response(buffer, filename):
    buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf',
                     as_attachment=True, download_name=filename)

def pdf_header(c, titulo, ancho, alto):
    from reportlab.lib.colors import HexColor
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # Fondo header
    c.setFillColor(HexColor('#0D1220'))
    c.rect(0, alto - 60, ancho, 60, fill=1, stroke=0)

    # Línea gradiente simulada con rectángulos
    colors = ['#4f8eff', '#3fa0f8', '#2fb8f0', '#1fcce8', '#0fe0e0', '#00e5a0']
    seg_w = ancho / len(colors)
    for i, col in enumerate(colors):
        c.setFillColor(HexColor(col))
        c.rect(i * seg_w, alto - 3, seg_w + 1, 3, fill=1, stroke=0)

    # Texto header
    c.setFillColor(HexColor('#4f8eff'))
    c.setFont('Helvetica-Bold', 13)
    c.drawString(30, alto - 38, nombre_negocio())

    c.setFillColor(HexColor('#7A92B8'))
    c.setFont('Helvetica', 9)
    c.drawString(30, alto - 52, titulo)

    c.setFillColor(HexColor('#3A5878'))
    c.setFont('Helvetica', 9)
    fecha = datetime.now().strftime('%d/%m/%Y')
    c.drawRightString(ancho - 30, alto - 38, fecha)

    return alto - 80

def pdf_tabla(c, y, headers, rows, ancho, col_widths=None):
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import Table, TableStyle
    from reportlab.lib import colors

    if not col_widths:
        col_w = (ancho - 60) / len(headers)
        col_widths = [col_w] * len(headers)

    data = [headers] + rows
    tabla = Table(data, colWidths=col_widths)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), HexColor('#1E2D45')),
        ('TEXTCOLOR', (0, 0), (-1, 0), HexColor('#7A92B8')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BACKGROUND', (0, 1), (-1, -1), HexColor('#0C1422')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [HexColor('#0C1422'), HexColor('#080B12')]),
        ('TEXTCOLOR', (0, 1), (-1, -1), HexColor('#C8D8F0')),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.3, HexColor('#1E2D45')),
        ('ROWHEIGHT', (0, 0), (-1, -1), 16),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    w, h = tabla.wrapOn(c, ancho - 60, 600)
    tabla.drawOn(c, 30, y - h)
    return y - h - 20

@export_bp.route('/api/export/pdf/inventario')
@login_required
def pdf_inventario():
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.colors import HexColor

    db = get_db()
    productos = db.execute('SELECT * FROM productos ORDER BY nombre').fetchall()
    db.close()

    buffer = io.BytesIO()
    ancho, alto = A4
    c = pdfcanvas.Canvas(buffer, pagesize=A4)
    c.setFillColor(HexColor('#080B12'))
    c.rect(0, 0, ancho, alto, fill=1, stroke=0)

    y = pdf_header(c, f"Inventario de productos — {datetime.now().strftime('%d/%m/%Y')}", ancho, alto)
    headers = ['Producto', 'Categoría', 'Stock', 'Mínimo', 'P. Venta', 'Estado']
    rows = []
    for p in productos:
        estado = 'Agotado' if p['stock'] == 0 else ('Bajo' if p['stock'] <= p['stock_minimo'] else 'OK')
        rows.append([p['nombre'], p['categoria'], str(p['stock']), str(p['stock_minimo']),
                     fmt_clp(p['precio_venta']), estado])
    col_widths = [140, 80, 45, 50, 70, 50]
    pdf_tabla(c, y, headers, rows, ancho, col_widths)
    c.save()
    return pdf_response(buffer, f"inventario_{datetime.now().strftime('%Y%m%d')}.pdf")

@export_bp.route('/api/export/pdf/ventas')
@login_required
def pdf_ventas():
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.colors import HexColor

    db = get_db()
    ventas = db.execute(
        'SELECT v.*, c.nombre as cn, u.nombre as un FROM ventas v '
        'LEFT JOIN clientes c ON v.cliente_id=c.id '
        'LEFT JOIN usuarios u ON v.usuario_id=u.id '
        'ORDER BY v.creado_en DESC LIMIT 100').fetchall()
    db.close()

    buffer = io.BytesIO()
    ancho, alto = A4
    c = pdfcanvas.Canvas(buffer, pagesize=A4)
    c.setFillColor(HexColor('#080B12'))
    c.rect(0, 0, ancho, alto, fill=1, stroke=0)

    y = pdf_header(c, f"Historial de ventas — {datetime.now().strftime('%d/%m/%Y')}", ancho, alto)
    headers = ['Fecha', 'Usuario', 'Cliente', 'Pago', 'Total']
    rows = []
    for v in ventas:
        dt = datetime.fromisoformat(v['creado_en'])
        rows.append([dt.strftime('%d/%m/%Y %H:%M'), v['un'] or '—',
                     v['cn'] or 'General', v['metodo_pago'], fmt_clp(v['total'])])
    col_widths = [100, 90, 110, 80, 65]
    pdf_tabla(c, y, headers, rows, ancho, col_widths)
    c.save()
    return pdf_response(buffer, f"ventas_{datetime.now().strftime('%Y%m%d')}.pdf")

@export_bp.route('/api/export/pdf/clientes')
@login_required
def pdf_clientes():
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.colors import HexColor

    db = get_db()
    clientes = db.execute('SELECT * FROM clientes ORDER BY nombre').fetchall()
    ventas_r = db.execute(
        'SELECT cliente_id, SUM(total) as total, COUNT(*) as nv FROM ventas GROUP BY cliente_id').fetchall()
    db.close()

    ventas_map = {v['cliente_id']: v for v in ventas_r}
    buffer = io.BytesIO()
    ancho, alto = A4
    c = pdfcanvas.Canvas(buffer, pagesize=A4)
    c.setFillColor(HexColor('#080B12'))
    c.rect(0, 0, ancho, alto, fill=1, stroke=0)

    y = pdf_header(c, f"Lista de clientes — {datetime.now().strftime('%d/%m/%Y')}", ancho, alto)
    headers = ['Nombre', 'Teléfono', 'Tipo', 'Compras', 'Total gastado']
    rows = []
    for cl in clientes:
        vd = ventas_map.get(cl['id'])
        rows.append([cl['nombre'], cl['telefono'] or '—', cl['tipo'],
                     str(vd['nv']) if vd else '0',
                     fmt_clp(vd['total']) if vd else '$0'])
    col_widths = [140, 100, 70, 60, 75]
    pdf_tabla(c, y, headers, rows, ancho, col_widths)
    c.save()
    return pdf_response(buffer, f"clientes_{datetime.now().strftime('%Y%m%d')}.pdf")

@export_bp.route('/api/export/pdf/stock-bajo')
@login_required
def pdf_stock_bajo():
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.colors import HexColor

    db = get_db()
    productos = db.execute(
        'SELECT * FROM productos WHERE stock <= stock_minimo ORDER BY stock ASC').fetchall()
    db.close()

    buffer = io.BytesIO()
    ancho, alto = A4
    c = pdfcanvas.Canvas(buffer, pagesize=A4)
    c.setFillColor(HexColor('#080B12'))
    c.rect(0, 0, ancho, alto, fill=1, stroke=0)

    y = pdf_header(c, f"Productos para reponer — {datetime.now().strftime('%d/%m/%Y')}", ancho, alto)
    headers = ['Producto', 'Categoría', 'Stock actual', 'Mínimo', 'Estado', 'P. Costo']
    rows = []
    for p in productos:
        estado = 'AGOTADO' if p['stock'] == 0 else 'BAJO'
        rows.append([p['nombre'], p['categoria'], str(p['stock']), str(p['stock_minimo']),
                     estado, fmt_clp(p['precio_costo'])])
    col_widths = [130, 80, 65, 50, 55, 65]
    pdf_tabla(c, y, headers, rows, ancho, col_widths)
    c.save()
    return pdf_response(buffer, f"stock_bajo_{datetime.now().strftime('%Y%m%d')}.pdf")
